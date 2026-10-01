"""
Parser de Excel de reservas de aulas.
Columnas: ÁREA SOLICITANTE | NÚM CC | NOMBRE PERSONA ABRE AULA | DÍA |
FECHA INI | HORA INI | FECHA FIN | HORA FIN | SEDE | AULA / CERRADURA
"""
from openpyxl import load_workbook
from typing import List, Dict, Any
import re
from datetime import datetime, time, date


def _norm(s: str) -> str:
    if s is None:
        return ""
    t = str(s).upper().strip()
    for a, b in [("Á", "A"), ("É", "E"), ("Í", "I"), ("Ó", "O"), ("Ú", "U"), ("Ñ", "N")]:
        t = t.replace(a, b)
    t = re.sub(r"[^A-Z0-9 /]+", " ", t)
    t = re.sub(r"\s+", " ", t).strip()
    return t


def _cell_str(v) -> str:
    if v is None:
        return ""
    if isinstance(v, datetime):
        return v.strftime("%d/%m/%Y")
    if isinstance(v, date):
        return v.strftime("%d/%m/%Y")
    if isinstance(v, time):
        return v.strftime("%H:%M:%S")
    return str(v).strip()


def _hora_str(v) -> str:
    if v is None:
        return ""
    if isinstance(v, datetime):
        return v.strftime("%H:%M:%S")
    if isinstance(v, time):
        return v.strftime("%H:%M:%S")
    if isinstance(v, float) and 0 <= v < 1:
        total = int(round(v * 24 * 3600))
        h, rem = divmod(total, 3600)
        m, sec = divmod(rem, 60)
        return f"{h:02d}:{m:02d}:{sec:02d}"
    return str(v).strip()


# Orden: más específico primero
HEADER_ALIASES = [
    ("aula", ["AULA / CERRADURA", "AULA/CERRADURA", "AULA CERRADURA", "CERRADURA"]),
    ("nombre", ["NOMBRE PERSONA ABRE AULA", "NOMBRE PERSONA", "PERSONA ABRE AULA"]),
    ("area", ["AREA SOLICITANTE", "AREA SOLICITANTE"]),
    ("cedula", ["NUM CC", "NÚM CC", "NUMERO CC", "NRO CC", "CEDULA"]),
    ("fecha_ini", ["FECHA INI", "FECHA INICIO", "FECHA INICIAL"]),
    ("hora_ini", ["HORA INI", "HORA INICIO", "HORA INICIAL"]),
    ("fecha_fin", ["FECHA FIN", "FECHA FINAL"]),
    ("hora_fin", ["HORA FIN", "HORA FINAL"]),
    ("sede", ["SEDE", "EDIFICIO"]),
    ("dia", ["DIA", "DÍA"]),
    # genéricos al final
    ("aula", ["AULA"]),
    ("nombre", ["NOMBRE"]),
    ("cedula", ["CC", "DOCUMENTO"]),
]


def parsear_reservas_desde_excel(ruta_archivo: str) -> List[Dict[str, Any]]:
    wb = load_workbook(ruta_archivo, data_only=True)
    ws = wb[wb.sheetnames[0]]

    header_row = None
    col_map = {}

    for r in range(1, min(40, (ws.max_row or 1) + 1)):
        cells = []
        for c in range(1, min(25, (ws.max_column or 1) + 1)):
            txt = _cell_str(ws.cell(r, c).value)
            if txt:
                cells.append((c, _norm(txt), txt))

        temp = {}
        used_cols = set()
        for field, aliases in HEADER_ALIASES:
            if field in temp:
                continue
            for alias in aliases:
                an = _norm(alias)
                for c, nn, raw in cells:
                    if c in used_cols:
                        continue
                    # coincidencia exacta o alias contenido en encabezado
                    if nn == an or nn.startswith(an + " ") or an == nn:
                        temp[field] = c
                        used_cols.add(c)
                        break
                    # "AULA / CERRADURA" contiene AULA y CERRADURA
                    if field == "aula" and "AULA" in nn and ("CERRADURA" in nn or nn == "AULA"):
                        # no tomar si dice NOMBRE PERSONA ABRE AULA
                        if "NOMBRE" in nn or "PERSONA" in nn:
                            continue
                        temp[field] = c
                        used_cols.add(c)
                        break
                if field in temp:
                    break

        if len(temp) >= 4 and "aula" in temp:
            header_row = r
            col_map = temp
            break

    if not header_row:
        wb.close()
        raise ValueError(
            "No se encontró encabezado de reservas. "
            "Columnas esperadas: ÁREA SOLICITANTE, NÚM CC, NOMBRE PERSONA ABRE AULA, "
            "DÍA, FECHA INI, HORA INI, FECHA FIN, HORA FIN, SEDE, AULA/CERRADURA"
        )

    reservas = []
    for r in range(header_row + 1, (ws.max_row or header_row) + 1):
        def get(field, hora=False):
            c = col_map.get(field)
            if not c:
                return ""
            v = ws.cell(r, c).value
            return _hora_str(v) if hora else _cell_str(v)

        aula = get("aula")
        if not aula:
            continue
        if _norm(aula) in ("AULA", "AULA / CERRADURA", "AULA CERRADURA"):
            continue

        reservas.append({
            "area_solicitante": get("area"),
            "cedula": get("cedula"),
            "nombre_persona": get("nombre"),
            "dia": get("dia"),
            "fecha_ini": get("fecha_ini"),
            "hora_ini": get("hora_ini", hora=True),
            "fecha_fin": get("fecha_fin"),
            "hora_fin": get("hora_fin", hora=True),
            "sede": get("sede"),
            "aula_codigo": aula.strip(),
        })

    wb.close()
    return reservas
